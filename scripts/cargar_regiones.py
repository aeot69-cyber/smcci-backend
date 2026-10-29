# -*- coding: utf-8 -*-
"""Carga regiones.xlsx (16 regiones / 56 provincias / 346 comunas) en la BD.

- Las 16 regiones ya existen: se enlazan por `codigo` (XV, I, ... RM).
- Las provincias se crean/actualizan por (región, nombre normalizado).
- Las comunas: si ya existe una con el mismo nombre normalizado se ACTUALIZA
  en el lugar (nombre oficial con tildes, región y provincia), conservando su
  id para no romper las visitas/hermanos que la usan. Si no existe, se crea.
- Idempotente: se puede volver a ejecutar sin duplicar nada.
- Todo se hace en UNA transacción (commit al final, rollback si algo falla).

Uso:  python cargar_regiones.py [ruta_regiones.xlsx]
"""
import sys, io, os, unicodedata
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

import openpyxl

XLSX = sys.argv[1] if len(sys.argv) > 1 else r"D:\Datos\MCCI\regiones.xlsx"
sys.path.insert(0, r"D:\Datos\MCCI\SMCCI_MULTI\smcci-backend")
os.environ.setdefault("DATABASE_URL",
    "postgresql://postgres.dcmwyuezcmukiwcxbpxp:Mcci2026SupabaseXq9vZt4kW7"
    "@aws-0-us-west-2.pooler.supabase.com:5432/postgres?sslmode=require")
import app as appmod

q = appmod.q


def norm(s):
    s = " ".join((s or "").strip().split())
    return "".join(c for c in unicodedata.normalize("NFD", s)
                   if unicodedata.category(c) != "Mn").lower()


def exec_sql(sql, p=()):
    """Ejecuta una sentencia sin esperar resultados (UPDATE/INSERT sin RETURNING)."""
    c = appmod.db()
    with c.cursor() as cur:
        cur.execute(sql, p)


def main():
    wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows or str(rows[0][0]).strip().lower() != "región":
        raise SystemExit("El archivo no tiene la cabecera esperada (Región, Identificador, Provincia, Comuna)")

    datos = []
    for reg, cod, prov, com in rows[1:]:
        if not (reg and cod and prov and com):
            continue
        datos.append((str(reg).strip(), str(cod).strip(), str(prov).strip(), str(com).strip()))
    print("Filas de datos en el Excel: %d" % len(datos))

    # Una sola conexión y UNA transacción para todo el proceso.
    with appmod.app.app_context():
        try:
            regiones = {norm(r["codigo"]): r for r in q("SELECT id, codigo, nombre FROM regiones")}
            if len(regiones) != 16:
                raise SystemExit("Se esperaban 16 regiones en la BD y hay %d" % len(regiones))

            # ---------- PROVINCIAS ----------
            prov_db = {(r["region_id"], norm(r["nombre"])): r
                       for r in q("SELECT id, nombre, region_id FROM provincias")}
            prov_ids = {}
            creadas_prov = act_prov = 0
            for reg_nombre, cod, prov_nombre, _ in datos:
                r = regiones.get(norm(cod))
                if not r:
                    raise SystemExit("Región del Excel no encontrada en la BD: %r (%s)" % (reg_nombre, cod))
                clave = (r["id"], norm(prov_nombre))
                if clave in prov_db:
                    pid = prov_db[clave]["id"]
                    if prov_db[clave]["nombre"] != prov_nombre:
                        exec_sql("UPDATE provincias SET nombre=%s WHERE id=%s", (prov_nombre, pid))
                        act_prov += 1
                else:
                    pid = q("INSERT INTO provincias(nombre, region_id) VALUES(%s,%s) RETURNING id",
                            (prov_nombre, r["id"]), one=True)["id"]
                    prov_db[clave] = {"id": pid, "nombre": prov_nombre, "region_id": r["id"]}
                    creadas_prov += 1
                prov_ids[(r["id"], norm(prov_nombre))] = pid
            print("Provincias: %d creadas, %d renombradas" % (creadas_prov, act_prov))

            # ---------- COMUNAS ----------
            com_db = {}
            for r in q("SELECT id, nombre, region_id, provincia_id FROM comunas"):
                com_db.setdefault(norm(r["nombre"]), []).append(r)

            nuevas = actualizadas = 0
            for reg_nombre, cod, prov_nombre, com_nombre in datos:
                rid = regiones[norm(cod)]["id"]
                pid = prov_ids[(rid, norm(prov_nombre))]
                existentes = com_db.get(norm(com_nombre), [])
                if existentes:
                    c = existentes[0]
                    if (c["nombre"], c["region_id"], c["provincia_id"]) != (com_nombre, rid, pid):
                        exec_sql("UPDATE comunas SET nombre=%s, region_id=%s, provincia_id=%s WHERE id=%s",
                                 (com_nombre, rid, pid, c["id"]))
                        actualizadas += 1
                else:
                    cid = q("INSERT INTO comunas(nombre, region_id, provincia_id) VALUES(%s,%s,%s) RETURNING id",
                            (com_nombre, rid, pid), one=True)["id"]
                    com_db.setdefault(norm(com_nombre), []).append(
                        {"id": cid, "nombre": com_nombre, "region_id": rid, "provincia_id": pid})
                    nuevas += 1
            print("Comunas: %d nuevas, %d actualizadas (en el lugar: conservan id)" % (nuevas, actualizadas))
            appmod.db().commit()
        except Exception:
            appmod.db().rollback()
            raise

        # ---------- VERIFICACIÓN ----------
        tot_prov = q("SELECT count(*) AS c FROM provincias", one=True)["c"]
        tot_com = q("SELECT count(*) AS c FROM comunas", one=True)["c"]
        sin_prov = q("SELECT count(*) AS c FROM comunas WHERE provincia_id IS NULL", one=True)["c"]
        print("BD: %d provincias, %d comunas, sin provincia: %d" % (tot_prov, tot_com, sin_prov))
        for fila in q("""SELECT r.nombre_corto AS region, p.nombre AS provincia, count(c.id) AS n
                           FROM provincias p JOIN regiones r ON r.id=p.region_id
                           LEFT JOIN comunas c ON c.provincia_id=p.id
                          GROUP BY r.nombre_corto, p.nombre ORDER BY r.nombre_corto, p.nombre"""):
            print("  %-14s %-28s %3d" % (fila["region"], fila["provincia"], fila["n"]))

        en_excel = {norm(x[3]) for x in datos}
        sobrantes = [r for r in q("SELECT id, nombre FROM comunas ORDER BY nombre")
                     if norm(r["nombre"]) not in en_excel]
        if sobrantes:
            print("AVISO: comunas en BD que no están en el Excel (no se modifican):",
                  [(s["id"], s["nombre"]) for s in sobrantes])
        else:
            print("Todas las comunas de la BD provienen del Excel ✔")


if __name__ == "__main__":
    main()
